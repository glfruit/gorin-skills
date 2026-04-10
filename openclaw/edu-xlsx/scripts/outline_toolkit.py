from __future__ import annotations

from collections import Counter
from pathlib import Path
import re
import unicodedata

import openpyxl
import pandas as pd

TEMPLATE_RULES = {
    "textbook": {
        "header_rows": (0, 4),
        "required_columns": ["项目名称", "任务名称", "模块/板块", "单元名称", "知识点/技能点"],
        "signatures": [
            {"项目名称", "任务名称", "模块/板块"},
            {"项目名称", "任务名称", "模块", "板块"},
        ],
        "knowledge_col": "知识点/技能点",
        "level_columns": ["项目名称", "任务名称", "模块/板块"],
        "group_column": "模块/板块",
        "content_column": "主要内容概要",
    },
    "course": {
        "header_rows": (0,),
        "required_columns": ["项目/模块名称", "任务/单元名称", "学时（理论）", "学时（实践）", "学时（总计）", "教学目标"],
        "signatures": [
            {"项目/模块名称", "任务/单元名称", "学时（总计）"},
            {"项目/模块名称", "任务/单元名称", "学时"},
        ],
        "knowledge_col": "教学目标",
        "level_columns": ["项目/模块名称", "任务/单元名称"],
        "group_column": "项目/模块名称",
        "content_column": "教学目标",
    },
    "training": {
        "header_rows": (0,),
        "required_columns": ["培训模块名称", "培训活动名称", "时长（分钟）", "培训对象/人数", "培训方式", "活动内容描述"],
        "signatures": [
            {"培训模块名称", "培训活动名称", "时长（分钟）"},
            {"培训模块名称", "培训活动名称", "时长"},
        ],
        "knowledge_col": "活动内容描述",
        "level_columns": ["培训模块名称", "培训活动名称"],
        "group_column": "培训模块名称",
        "content_column": "活动内容描述",
    },
}

PLACEHOLDER_PATTERN = re.compile(r"(?:\bTBD\b|\bTODO\b|待定|适当|酌情)", re.IGNORECASE)
TERM_SPLIT_PATTERN = re.compile(r"[、,，；;\s/]+")


def _normalize_text(value: object) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value))
    return re.sub(r"\s+", "", text).strip()


def _display_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _load_sheet(path: str | Path, read_only: bool = True):
    wb = openpyxl.load_workbook(path, read_only=read_only, data_only=False)
    return wb, wb.active


def _row_cells(ws, row_number: int):
    rows = ws.iter_rows(min_row=row_number, max_row=row_number)
    return next(rows, [])


def _headers_for_row(ws, row_number: int) -> list[str]:
    return [_display_text(cell.value) for cell in _row_cells(ws, row_number)]


def _normalized_headers(ws, row_number: int) -> set[str]:
    return {_normalize_text(cell.value) for cell in _row_cells(ws, row_number) if _normalize_text(cell.value)}


def identify(path: str | Path) -> str | None:
    wb, ws = _load_sheet(path, read_only=True)
    try:
        rows_to_check = [1, 5]
        best_type = None
        best_score = 0
        for row_number in rows_to_check:
            headers = _normalized_headers(ws, row_number)
            if not headers:
                continue
            for template_type, rule in TEMPLATE_RULES.items():
                score = max(
                    len({_normalize_text(item) for item in signature} & headers)
                    for signature in rule["signatures"]
                )
                if score > best_score:
                    best_score = score
                    best_type = template_type
        return best_type if best_score >= 3 else None
    finally:
        wb.close()


def read_outline(path: str | Path, template_type: str | None = None) -> pd.DataFrame:
    template_type = template_type or identify(path)
    if template_type not in TEMPLATE_RULES:
        raise ValueError(f"Unable to identify outline template for: {path}")

    preferred_header = 4 if template_type == "textbook" else 0
    candidate_headers = [preferred_header] if preferred_header == 0 else [preferred_header, 0]

    last_error: Exception | None = None
    df: pd.DataFrame | None = None
    for header in candidate_headers:
        try:
            df = pd.read_excel(path, header=header)
            if any(_normalize_text(col) for col in df.columns):
                break
        except ValueError as exc:
            last_error = exc
            continue

    if df is None:
        raise ValueError(f"Unable to read outline: {path}") from last_error

    cleaned_columns = []
    for idx, col in enumerate(df.columns):
        text = _display_text(col)
        cleaned_columns.append(text if text else f"Unnamed:{idx}")
    df.columns = cleaned_columns
    return df


def _resolve_template_type(path: str | Path, template_type: str | None) -> str:
    resolved = template_type or identify(path)
    if resolved not in TEMPLATE_RULES:
        raise ValueError(f"Unable to identify outline template for: {path}")
    return resolved


def _nonempty_row_mask(df: pd.DataFrame) -> pd.Series:
    normalized = df.fillna("").astype(str)
    normalized = normalized.apply(lambda column: column.map(lambda value: value.strip()))
    return normalized.ne("").any(axis=1)


def _detect_header_row(path: str | Path, template_type: str) -> int:
    wb, ws = _load_sheet(path, read_only=True)
    try:
        for candidate in (1, 5):
            headers = _normalized_headers(ws, candidate)
            if not headers:
                continue
            for signature in TEMPLATE_RULES[template_type]["signatures"]:
                if len({_normalize_text(item) for item in signature} & headers) >= 3:
                    return candidate - 1
        return 4 if template_type == "textbook" else 0
    finally:
        wb.close()


def _scan_placeholders(df: pd.DataFrame, nonempty_mask: pd.Series) -> list[str]:
    warnings: list[str] = []
    for row_index, row in df[nonempty_mask].iterrows():
        for column, value in row.items():
            text = _display_text(value)
            if text and PLACEHOLDER_PATTERN.search(text):
                warnings.append(f"第{row_index + 2}行 {column} 包含占位符: {text}")
    return warnings


def _check_hierarchy(df: pd.DataFrame, template_type: str, nonempty_mask: pd.Series) -> list[str]:
    errors: list[str] = []
    columns = TEMPLATE_RULES[template_type]["level_columns"]
    for row_index, row in df[nonempty_mask].iterrows():
        values = [_display_text(row.get(column)) for column in columns]
        for idx in range(1, len(values)):
            if values[idx] and not all(values[:idx]):
                parent = columns[idx - 1]
                current = columns[idx]
                errors.append(f"第{row_index + 2}行层级跳级: {current} 有值但 {parent} 为空")
    return errors


def _check_knowledge_gaps(df: pd.DataFrame, template_type: str, nonempty_mask: pd.Series) -> list[str]:
    column = TEMPLATE_RULES[template_type]["knowledge_col"]
    if column not in df.columns:
        return [f"缺少关键列: {column}"]

    errors: list[str] = []
    for row_index, row in df[nonempty_mask].iterrows():
        if not _display_text(row.get(column)):
            errors.append(f"第{row_index + 2}行 {column} 为空")
    return errors


def _check_required_columns(df: pd.DataFrame, template_type: str) -> list[str]:
    columns = {_normalize_text(column): column for column in df.columns}
    missing = []
    for required in TEMPLATE_RULES[template_type]["required_columns"]:
        if _normalize_text(required) not in columns:
            missing.append(required)
    return missing


def _check_terminology(df: pd.DataFrame, template_type: str, nonempty_mask: pd.Series) -> list[str]:
    group_column = TEMPLATE_RULES[template_type]["group_column"]
    if group_column not in df.columns:
        return []

    terms = Counter()
    examples: dict[str, str] = {}
    for _, row in df[nonempty_mask].iterrows():
        text = _display_text(row.get(group_column))
        if not text:
            continue
        for token in TERM_SPLIT_PATTERN.split(text):
            token = token.strip()
            if len(token) < 2:
                continue
            key = token.lower()
            terms[key] += 1
            examples.setdefault(key, token)

    warnings = []
    for key, count in terms.items():
        example = examples[key]
        if len(example) >= 2 and terms.get(key.replace("模块", "板块"), 0) and "模块" in key:
            warnings.append(f"术语可能不一致: 同时出现“{example}”与“{example.replace('模块', '板块')}”")
        if len(example) >= 2 and terms.get(key.replace("板块", "模块"), 0) and "板块" in key:
            warnings.append(f"术语可能不一致: 同时出现“{example.replace('板块', '模块')}”与“{example}”")
    return sorted(set(warnings))


def validate(path: str | Path, template_type: str | None = None, *, check_terms: bool = False) -> dict:
    template_type = _resolve_template_type(path, template_type)
    detected_header = _detect_header_row(path, template_type)
    df = read_outline(path, template_type)

    errors: list[str] = []
    warnings: list[str] = []

    missing_columns = _check_required_columns(df, template_type)
    if missing_columns:
        errors.append(f"缺少必填列: {', '.join(missing_columns)}")

    nonempty_mask = _nonempty_row_mask(df)
    errors.extend(_check_knowledge_gaps(df, template_type, nonempty_mask))
    errors.extend(_check_hierarchy(df, template_type, nonempty_mask))
    warnings.extend(_scan_placeholders(df, nonempty_mask))
    if check_terms:
        warnings.extend(_check_terminology(df, template_type, nonempty_mask))

    stats = {
        "template_type": template_type,
        "header_row": detected_header + 1,
        "total_rows": int(len(df)),
        "nonempty_rows": int(nonempty_mask.sum()),
        "error_count": len(errors),
        "warning_count": len(warnings),
    }
    return {
        "passed": not errors,
        "errors": errors,
        "warnings": warnings,
        "stats": stats,
    }


def extract_sections(path: str | Path, template_type: str | None = None) -> str:
    template_type = _resolve_template_type(path, template_type)
    df = read_outline(path, template_type)
    nonempty_mask = _nonempty_row_mask(df)

    rule = TEMPLATE_RULES[template_type]
    group_column = rule["group_column"]
    knowledge_column = rule["knowledge_col"]
    content_column = rule["content_column"]

    sections: list[str] = []
    current_group = "未分组"
    for row_index, row in df[nonempty_mask].iterrows():
        group_value = _display_text(row.get(group_column))
        if group_value:
            current_group = group_value

        knowledge_value = _display_text(row.get(knowledge_column)) or "（未填写）"
        content_value = _display_text(row.get(content_column)) or knowledge_value
        sections.append(
            "\n".join(
                [
                    f"## 板块: {current_group}",
                    "",
                    f"### 知识点: {knowledge_value}",
                    f"- **内容**: {content_value}",
                    f"- **来源**: 大纲第{row_index + 2}行",
                ]
            )
        )

    return "\n\n".join(sections).strip()
